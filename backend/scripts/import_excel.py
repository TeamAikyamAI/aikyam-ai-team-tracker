"""Import the team's Weekly_Update.xlsx into the tracker.

    cd backend
    python scripts/import_excel.py "C:/path/Weekly_Update.xlsx"            # do it
    python scripts/import_excel.py "C:/path/Weekly_Update.xlsx" --dry-run  # show what would happen

Workbook layout expected - the same one the tracker's own export produces
(Dashboard > Export), one sheet, first row = header:

    Project Name | Assigned by | Date of assignment | Status |
    Week update - DD.MM.YYYY ... | Date of completion | Remarks

Any "Week ... DD.MM.YYYY" column becomes a dated update against the project,
so a workbook exported from the tracker can be edited by the team and read back
in without losing the week it belonged to. "Project Description" is still read
if the file has one.

Older files with one sheet per pipeline stage still work: when the workbook has
more than one sheet, the sheet name is the status, as before.

What it does, idempotently (safe to run twice):
  * files every project under the status its own Status cell names, creating
    that status if the tracker does not have it yet - whatever it is called. A
    blank Status cell lands in the first status of the pipeline.
  * a status created here is never marked terminal: which stage means
    "delivered" is a decision for Admin > Statuses, and guessing it wrong marks
    every imported project as finished.
  * creates the "--vertical" (default: Unassigned) if missing, so every row
    lands somewhere and can be re-homed from Admin > Verticals later
  * skips rows whose project name already exists (case-insensitive) and
    duplicate names inside the workbook itself
"""
from __future__ import annotations

import argparse
import os
import sys
from datetime import date, datetime, time, timezone
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))
os.chdir(BACKEND_DIR)

from dotenv import load_dotenv  # noqa: E402

load_dotenv(BACKEND_DIR / ".env")

from openpyxl import load_workbook  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
import app.models  # noqa: E402,F401
from app.models.project import Project  # noqa: E402
from app.models.status import Status  # noqa: E402
from app.models.update import Update  # noqa: E402
from app.models.user import User  # noqa: E402
from app.models.vertical import Vertical  # noqa: E402
from app.services import daily  # noqa: E402
from app.services import weeks as wk  # noqa: E402
from app.services.audit import log_action  # noqa: E402
from app.services.excel import text_to_update  # noqa: E402
from sqlalchemy import func  # noqa: E402

HEADER_ALIASES = {
    "project name": "name",
    "project description": "description",
    "assigned by": "assigned_by",
    "date of assignment": "assigned_on",
    "status": "status_note",
    "date of completion": "actual_completion_date",
    "remarks": "remarks",
}

# Row-level status spellings that mean "same as the sheet", so they are not
# copied into remarks as noise. Only used for the old one-sheet-per-stage files.
STATUS_SYNONYMS = {
    "wip": {"wip", "in progress", "in-progress", "in-progess", "inprogress", "ongoing"},
    "uat": {"uat", "in uat", "in-uat", "testing"},
    "live": {"live", "completed", "done", "in production"},
    "in queue": {"in queue", "queue", "queued", "pending", "backlog"},
}

# A hint, not a list of allowed statuses: any status name in the file is created,
# and one that is not named here simply gets grey.
STAGE_COLOURS = {
    "in queue": "#94a3b8",
    "wip": "#3b82f6",
    "uat": "#f59e0b",
    "live": "#22c55e",
    "on hold": "#f97316",
    "blocked": "#ef4444",
}

PLACEHOLDER_STATUSES = ("Ongoing", "Blocked", "On Hold", "Completed")

# What people write in a week cell when there is nothing to report.
EMPTY_CELL_VALUES = {"na", "n/a", "-", "--", "nil", "none"}

DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d.%m.%y",
                "%Y/%m/%d", "%Y.%m.%d", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y")


def _norm(value) -> str:
    return " ".join(str(value or "").split()).strip().lower()


def _as_date(value, where: str = "") -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = " ".join(str(value).split()).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    # Name the row: in a 200-row workbook "could not parse '24-.8.2026'" on its
    # own tells nobody which project lost its date.
    print(f"  ! {where or 'a row'}: could not read the date {value!r} - left blank, fix it in the tracker")
    return None


def _clean(value) -> str | None:
    text = " ".join(str(value).split()).strip() if value is not None else ""
    return text or None


def _cell_text(value) -> str | None:
    """A week cell's content, or None when it says nothing."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in EMPTY_CELL_VALUES:
        return None
    return text


def read_rows(path: Path) -> list[dict]:
    """Every project row in the workbook, with its week updates attached."""
    wb = load_workbook(path, data_only=True)
    rows: list[dict] = []
    for order, ws in enumerate(wb.worksheets):
        header = [(_norm(c)) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        mapping = {i: HEADER_ALIASES[h] for i, h in enumerate(header) if h in HEADER_ALIASES}
        if "name" not in mapping.values():
            print(f"  sheet {ws.title!r}: no 'Project Name' column, skipped")
            continue
        # Week columns are read off the headings, never off a stored list - the
        # same function the export writes them with reads them back.
        week_columns = {i: start for i, cell in enumerate(header)
                        if (start := wk.parse_header(cell)) is not None}
        for row_no, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            record = {field: raw[i] if i < len(raw) else None for i, field in mapping.items()}
            name = _clean(record.get("name"))
            if not name:
                continue
            where = f"{ws.title!r} row {row_no} ({name})"
            updates = [
                (start, text)
                for i, start in sorted(week_columns.items(), key=lambda kv: kv[1])
                if i < len(raw) and (text := _cell_text(raw[i]))
            ]
            rows.append({
                "sheet": ws.title.strip(),
                "sheet_order": order,
                "row_no": row_no,
                "where": where,
                "name": name,
                "description": _clean(record.get("description")),
                "assigned_by": _clean(record.get("assigned_by")),
                "assigned_on": _as_date(record.get("assigned_on"), f"{where}, Date of assignment"),
                "status_note": _clean(record.get("status_note")),
                "actual_completion_date": _as_date(record.get("actual_completion_date"),
                                                   f"{where}, Date of completion"),
                "remarks": _clean(record.get("remarks")),
                "updates": updates,
            })
    return rows


def _status_names(rows: list[dict], single_sheet: bool) -> list[str]:
    """The statuses the workbook asks for, in the order it first mentions them.

    One sheet means the Status column is the status. Several sheets is the old
    layout, where the sheet name was.
    """
    names: list[str] = []
    seen: set[str] = set()
    source = (r["status_note"] for r in rows) if single_sheet else (r["sheet"] for r in rows)
    for value in source:
        name = _clean(value)
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        names.append(name)
    return names


def _first_in_pipeline(db, pending: list[Status]) -> Status | None:
    """The first stage by sort order - where a row with no Status belongs."""
    candidates = list(db.query(Status).filter(Status.is_active == True).all())  # noqa: E712
    candidates += [s for s in pending if s.id is None]  # dry run: never flushed
    if not candidates:
        return None
    return sorted(candidates, key=lambda s: (s.sort_order or 0, s.id if s.id is not None else 10 ** 9))[0]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", help="Path to Weekly_Update.xlsx")
    ap.add_argument("--vertical", default="Unassigned", help="Vertical to file imported projects under (created if missing)")
    ap.add_argument("--created-by", default=None, help="E-mail of the tracker user to record as creator (default: first active admin)")
    ap.add_argument("--keep-placeholders", action="store_true", help="Do not deactivate unused seed statuses the workbook never mentions")
    ap.add_argument("--dry-run", action="store_true", help="Print the plan, change nothing")
    args = ap.parse_args()

    path = Path(args.workbook).expanduser()
    if not path.is_file():
        print(f"Workbook not found: {path}")
        return 2

    rows = read_rows(path)
    sheets = sorted({(r["sheet_order"], r["sheet"]) for r in rows})
    single_sheet = len(sheets) == 1
    print(f"Read {len(rows)} row(s) from {path.name} across {len(sheets)} sheet(s)"
          + (" - statuses come from the Status column" if single_sheet else " - statuses come from the sheet names"))

    db = SessionLocal()
    try:
        creator = None
        if args.created_by:
            creator = db.query(User).filter(func.lower(User.email) == args.created_by.lower()).first()
            if not creator:
                print(f"No user with e-mail {args.created_by}")
                return 2
        else:
            creator = db.query(User).filter(User.role == "admin", User.is_active == True).order_by(User.id).first()  # noqa: E712
        if not creator:
            print("No active admin found - run scripts/seed.py first")
            return 2
        print(f"Recording {creator.name} <{creator.email}> as creator")

        # --- vertical -------------------------------------------------------
        vertical = db.query(Vertical).filter(func.lower(Vertical.name) == args.vertical.lower()).first()
        if not vertical:
            print(f"Vertical {args.vertical!r}: will create (head = {creator.name})")
            vertical = Vertical(name=args.vertical, head_name=creator.name, head_email=creator.email, is_active=True)
            if not args.dry_run:
                db.add(vertical)
                db.flush()
        else:
            print(f"Vertical {vertical.name!r}: exists")

        # --- statuses the workbook asks for ---------------------------------
        wanted = _status_names(rows, single_sheet)
        next_order = (db.query(func.max(Status.sort_order)).scalar() or 0) + 1
        status_by_name: dict[str, Status] = {}
        pending: list[Status] = []
        for idx, name in enumerate(wanted):
            st = db.query(Status).filter(func.lower(Status.name) == name.lower()).first()
            if st:
                if not st.is_active:
                    print(f"Status {name!r}: exists but inactive -> re-activating")
                    st.is_active = True
                else:
                    print(f"Status {name!r}: exists")
            else:
                # Never terminal: only an admin knows which stage means delivered,
                # and getting it wrong stamps a completion date on every project.
                order = next_order + idx
                print(f"Status {name!r}: will create (order {order}) - tick 'terminal' in Admin > Statuses if this is the finished stage")
                st = Status(name=name, color=STAGE_COLOURS.get(name.lower(), "#6b7280"),
                            is_terminal=False, sort_order=order, is_active=True)
                pending.append(st)
                if not args.dry_run:
                    db.add(st)
                    db.flush()
            status_by_name[name.lower()] = st

        fallback = _first_in_pipeline(db, pending)
        if fallback is None:
            print("No statuses exist and the workbook names none - add one in Admin > Statuses first")
            return 2

        if not args.keep_placeholders:
            named = {n.lower() for n in wanted}
            for name in PLACEHOLDER_STATUSES:
                st = db.query(Status).filter(Status.name == name, Status.is_active == True).first()  # noqa: E712
                # A status the workbook uses is a real status, whether or not it
                # happens to share a name with a seed placeholder.
                if not st or name.lower() in named:
                    continue
                in_use = db.query(Project.id).filter(Project.status_id == st.id).first() is not None
                if in_use:
                    print(f"Placeholder status {name!r}: in use, left active")
                else:
                    print(f"Placeholder status {name!r}: unused and not in the workbook -> deactivating (re-enable any time in Admin > Statuses)")
                    if not args.dry_run:
                        st.is_active = False

        # --- projects ---------------------------------------------------------
        tz = daily.tzinfo_for(db)
        existing = {n.lower() for (n,) in db.query(Project.name).all()}
        seen_in_file: set[str] = set()
        created = skipped = logged = 0
        for r in rows:
            key = r["name"].lower()
            if key in existing:
                print(f"  - {r['name']!r}: already in tracker, skipped")
                skipped += 1
                continue
            if key in seen_in_file:
                print(f"  - {r['name']!r}: duplicate row in workbook, skipped")
                skipped += 1
                continue
            seen_in_file.add(key)

            remarks = r["remarks"]
            note = _norm(r["status_note"])
            if single_sheet:
                st = status_by_name.get(note, fallback)
            else:
                sheet_key = _norm(r["sheet"])
                st = status_by_name[sheet_key]
                if note and note != sheet_key and note not in STATUS_SYNONYMS.get(sheet_key, set()):
                    line = f"Status note from import: {r['status_note']}"
                    remarks = f"{remarks}\n{line}" if remarks else line

            completion = r["actual_completion_date"]
            print(f"  + {r['name']!r} -> {st.name}"
                  + (f", assigned by {r['assigned_by']}" if r["assigned_by"] else "")
                  + (f" on {r['assigned_on']}" if r["assigned_on"] else "")
                  + (f", completed {completion}" if completion else "")
                  + (f", {len(r['updates'])} week update(s)" if r["updates"] else "")
                  + (f", remarks: {remarks!r}" if remarks else ""))
            if not args.dry_run:
                project = Project(
                    name=r["name"],
                    description=r["description"],
                    vertical_id=vertical.id,
                    status_id=st.id,
                    created_by_id=creator.id,
                    assigned_by=r["assigned_by"],
                    assigned_on=r["assigned_on"],
                    remarks=remarks,
                    actual_completion_date=completion,
                )
                project.owners = []
                db.add(project)
                db.flush()
                for week_start, text in r["updates"]:
                    # Stamped at midday of the week it was reported in, so it
                    # falls back into the same column on the next export whatever
                    # the timezone does at the boundaries.
                    db.add(Update(
                        project_id=project.id,
                        author_id=creator.id,
                        created_at=datetime.combine(week_start, time(12, 0), tzinfo=tz).astimezone(timezone.utc),
                        **text_to_update(text),
                    ))
            logged += len(r["updates"])
            created += 1

        if args.dry_run:
            db.rollback()
            print(f"\nDry run: {created} project(s) would be created with {logged} update(s), "
                  f"{skipped} skipped. Nothing changed.")
        else:
            db.commit()
            log_action(db, creator.id, "import_excel",
                       details=f"{path.name}: {created} created, {logged} updates, {skipped} skipped")
            print(f"\nDone: {created} project(s) created with {logged} update(s), {skipped} skipped."
                  f"\nNext: open the dashboard, set owners on each project, and move any project"
                  f" from '{vertical.name}' to its real vertical (Edit project > Vertical).")
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
