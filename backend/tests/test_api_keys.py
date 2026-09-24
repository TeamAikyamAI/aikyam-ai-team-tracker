"""The API key register.

A record of what exists and what is about to lapse - never of the secrets
themselves. These tests hold both halves of that: the register works, and there
is nowhere in it for a key value to live.
"""
from datetime import date, timedelta

import pytest

TODAY = date.today()


@pytest.fixture()
def provider(client, auth):
    rows = client.get("/api-providers", headers=auth("admin")).json()
    if rows:
        return rows[0]["id"]
    return client.post("/api-providers", headers=auth("admin"), json={"name": "Groq"}).json()["id"]


def _entry(client, auth, provider, who="admin", **over):
    body = {
        "provider_id": provider,
        "project_label": "Bulk email",
        "purpose": "For bulk emailing",
        "status": "active",
    }
    body.update(over)
    return client.post("/api-keys", headers=auth(who), json=body)


# ---------------------------------------------------------------------------
# who can see it
# ---------------------------------------------------------------------------

def test_the_register_is_for_the_ai_team_only(client, seed, auth):
    """Requestors have no reason to see the team's infrastructure."""
    assert client.get("/api-keys", headers=auth("admin")).status_code == 200
    assert client.get("/api-keys", headers=auth("member")).status_code == 200
    assert client.get("/api-keys", headers=auth("req1")).status_code == 403
    assert client.get("/api-providers", headers=auth("req1")).status_code == 403
    assert client.get("/api-keys/export.xlsx", headers=auth("req1")).status_code == 403


def test_a_requestor_cannot_write_to_it_either(client, seed, auth, provider):
    assert _entry(client, auth, provider, who="req1").status_code == 403


def test_the_feature_can_be_handed_to_a_role_from_the_grid(client, seed, auth, db):
    """Access is a grid toggle like everything else, not a hardcoded role."""
    from app.services import permissions as perms
    perms.set_many(db, {"api_keys": {"requestor": True}})
    try:
        assert client.get("/api-keys", headers=auth("req1")).status_code == 200
    finally:
        perms.set_many(db, {"api_keys": {"requestor": False}})
    assert client.get("/api-keys", headers=auth("req1")).status_code == 403


# ---------------------------------------------------------------------------
# the register itself
# ---------------------------------------------------------------------------

def test_an_entry_belongs_either_to_a_project_or_to_a_typed_label(client, seed, auth, provider):
    """Plenty of keys belong to work that never became a tracker project."""
    linked = _entry(client, auth, provider, project_id=None, project_label="FD_Rate")
    assert linked.status_code == 201
    assert linked.json()["project_name"] == "FD_Rate"

    made = client.post("/projects", headers=auth("admin"), json={
        "name": "Key Owning Project", "vertical_id": seed["vertical"], "status_id": seed["status_queue"],
    }).json()
    to_project = _entry(client, auth, provider, project_id=made["id"], project_label="ignored")
    assert to_project.status_code == 201
    # A linked project IS the name - a label alongside it would only drift.
    assert to_project.json()["project_name"] == "Key Owning Project"
    assert to_project.json()["project_label"] is None


def test_an_entry_with_neither_is_refused(client, seed, auth, provider):
    r = client.post("/api-keys", headers=auth("admin"), json={
        "provider_id": provider, "purpose": "orphan",
    })
    assert r.status_code == 422


def test_there_is_nowhere_to_put_the_secret(client, seed, auth, provider):
    """The register records that a key exists, not what it is. An attempt to
    smuggle a value in is dropped rather than stored."""
    r = _entry(client, auth, provider, key_value="sk-live-do-not-store-me",
               api_key="sk-live-do-not-store-me")
    assert r.status_code == 201
    body = r.json()
    assert "key_value" not in body and "api_key" not in body
    assert "sk-live" not in str(body)


def test_expiry_state_is_worked_out_not_stored(client, seed, auth, provider):
    """A row cannot claim to be fine on a date that has already passed."""
    gone = _entry(client, auth, provider, expires_on=str(TODAY - timedelta(days=3))).json()
    soon = _entry(client, auth, provider, expires_on=str(TODAY + timedelta(days=10))).json()
    far = _entry(client, auth, provider, expires_on=str(TODAY + timedelta(days=200))).json()
    never = _entry(client, auth, provider, expires_on=None).json()

    assert gone["expiry_state"] == "expired" and gone["days_left"] == -3
    assert soon["expiry_state"] == "soon" and soon["days_left"] == 10
    assert far["expiry_state"] == "ok"
    assert never["expiry_state"] == "none" and never["days_left"] is None


def test_a_revoked_key_stops_being_a_warning(client, seed, auth, provider):
    """Nobody needs chasing about a key that has already been turned off."""
    made = _entry(client, auth, provider, expires_on=str(TODAY - timedelta(days=5))).json()
    assert made["expiry_state"] == "expired"

    revoked = client.patch(f"/api-keys/{made['id']}", headers=auth("admin"),
                           json={"status": "revoked"}).json()
    assert revoked["expiry_state"] == "none"

    ids = [r["id"] for r in client.get("/api-keys/expiring", headers=auth("admin")).json()]
    assert made["id"] not in ids


def test_the_expiring_list_is_what_the_card_and_the_digest_count(client, seed, auth, provider):
    _entry(client, auth, provider, expires_on=str(TODAY + timedelta(days=2)), project_label="Soon one")
    _entry(client, auth, provider, expires_on=str(TODAY + timedelta(days=400)), project_label="Far one")

    rows = client.get("/api-keys/expiring", headers=auth("admin")).json()
    labels = [r["project_name"] for r in rows]
    assert "Soon one" in labels
    assert "Far one" not in labels
    assert all(r["expiry_state"] in ("soon", "expired") for r in rows)


def test_entries_can_be_edited_and_deleted(client, seed, auth, provider):
    made = _entry(client, auth, provider, purpose="First wording").json()
    edited = client.patch(f"/api-keys/{made['id']}", headers=auth("member"),
                          json={"purpose": "Better wording"})
    assert edited.status_code == 200
    assert edited.json()["purpose"] == "Better wording"

    assert client.delete(f"/api-keys/{made['id']}", headers=auth("member")).status_code == 204
    ids = [r["id"] for r in client.get("/api-keys", headers=auth("admin")).json()]
    assert made["id"] not in ids


# ---------------------------------------------------------------------------
# the provider master list
# ---------------------------------------------------------------------------

def test_providers_are_admin_managed(client, seed, auth):
    """Like verticals and statuses - a list, not free text, so "Groq" and
    "groq" never become two things."""
    made = client.post("/api-providers", headers=auth("admin"), json={"name": "Deep-infra"})
    assert made.status_code == 201

    dupe = client.post("/api-providers", headers=auth("admin"), json={"name": "deep-INFRA"})
    assert dupe.status_code == 400

    # A member reads the list to fill the form, but does not curate it.
    assert client.get("/api-providers", headers=auth("member")).status_code == 200
    assert client.post("/api-providers", headers=auth("member"), json={"name": "Nope"}).status_code == 403


def test_a_provider_in_use_cannot_be_deleted(client, seed, auth):
    made = client.post("/api-providers", headers=auth("admin"), json={"name": "InUseProvider"}).json()
    _entry(client, auth, made["id"], project_label="Something")

    refused = client.delete(f"/api-providers/{made['id']}", headers=auth("admin"))
    assert refused.status_code == 400
    assert "Deactivate" in refused.json()["detail"]


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------

def test_the_export_matches_the_teams_own_columns(client, seed, auth, provider):
    import io
    from openpyxl import load_workbook

    _entry(client, auth, provider, project_label="FD_Rate", purpose="For the Scrapping",
           account_email="fd_treasury@aikyamcap.com", expires_on=None)

    r = client.get("/api-keys/export.xlsx", headers=auth("member"))
    assert r.status_code == 200
    assert ".xlsx" in r.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header[:5] == ["Project", "API-Key", "Purpose", "Exp-date", "email_id"]

    body = "\n".join(
        " | ".join("" if c is None else str(c) for c in row)
        for row in ws.iter_rows(min_row=2, values_only=True)
    )
    assert "FD_Rate" in body
    assert "No Expiry Date" in body          # a blank date is a real answer
    assert "fd_treasury@aikyamcap.com" in body


def test_the_weekly_rollup_raises_keys_that_are_about_to_lapse(client, seed, auth, provider, outbox):
    """The rollup is the one message that reaches the person who can get a
    renewal approved."""
    client.patch("/settings", headers=auth("admin"),
                 json={"values": {"smtp_host": "smtp.example.com"}})
    _entry(client, auth, provider, project_label="Meeting Hub",
           expires_on=str(TODAY + timedelta(days=5)), purpose="For the STT")
    outbox.clear()

    r = client.post("/settings/digest-preview", headers=auth("admin"))
    assert r.status_code == 200

    rollups = [m for m in outbox if "rollup" in m["subject"].lower()]
    assert rollups, "no rollup was produced"
    assert "API keys needing attention" in rollups[0]["html"]
    assert "Meeting Hub" in rollups[0]["html"]


# ---------------------------------------------------------------------------
# a key that exists on paper before it exists at the vendor
# ---------------------------------------------------------------------------

def test_a_pending_key_may_have_no_provider_yet(client, seed, auth):
    """The team's own sheet already records these: the work is planned, the
    key is not taken out, and which vendor it comes from is undecided. That
    row is exactly what the register is for - it must not be excluded."""
    r = client.post("/api-keys", headers=auth("admin"), json={
        "project_label": "Compliance Tracker",
        "purpose": "For the Ai agent",
        "status": "pending",
    })
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["provider_id"] is None
    assert body["provider_name"] == "Pending"


def test_a_live_key_must_say_where_it_came_from(client, seed, auth):
    refused = client.post("/api-keys", headers=auth("admin"), json={
        "project_label": "Something live", "status": "active",
    })
    assert refused.status_code == 422


def test_a_pending_key_cannot_be_switched_on_without_a_provider(client, seed, auth):
    made = client.post("/api-keys", headers=auth("admin"), json={
        "project_label": "Still deciding", "status": "pending",
    }).json()
    refused = client.patch(f"/api-keys/{made['id']}", headers=auth("admin"),
                           json={"status": "active"})
    assert refused.status_code == 400
    assert "provider" in refused.json()["detail"].lower()


def test_the_export_writes_pending_where_no_provider_is_known(client, seed, auth):
    import io
    from openpyxl import load_workbook

    client.post("/api-keys", headers=auth("admin"), json={
        "project_label": "Waiting on a decision", "purpose": "For the Ai agent",
        "status": "pending",
    })
    ws = load_workbook(io.BytesIO(client.get("/api-keys/export.xlsx", headers=auth("admin")).content)).active
    body = [
        " | ".join("" if c is None else str(c) for c in row)
        for row in ws.iter_rows(min_row=2, values_only=True)
    ]
    line = next(l for l in body if "Waiting on a decision" in l)
    assert "Pending" in line


# ---------------------------------------------------------------------------
# bringing the team's existing spreadsheet in
# ---------------------------------------------------------------------------

def test_the_existing_spreadsheet_can_be_imported(client, seed, auth, db, tmp_path):
    """An empty register is worth nothing - the point is to see what already
    exists. The import reads the sheet in the team's own column order."""
    from openpyxl import Workbook
    from scripts.import_api_keys import run

    project = client.post("/projects", headers=auth("admin"), json={
        "name": "Imported Project", "vertical_id": seed["vertical"], "status_id": seed["status_queue"],
    }).json()

    wb = Workbook()
    ws = wb.active
    ws.append(["Project", "API-Key", "Purpose", "Exp-date", "email_id"])
    ws.append(["FD_Rate_Import", "Gemini", "For the imported Scrapping", "No Expiry Date", "fd_treasury@aikyamcap.com"])
    ws.append(["Imported Project", "Groq", "For the imported fallback", None, None])
    ws.append(["Imported Project", "Pending", "For the imported Ai agent", None, None])
    ws.append(["Bulk email", "Brand-New-Vendor", "For the import test mailing", None, None])
    path = tmp_path / "Api_Key.xlsx"
    wb.save(path)

    import app.core.database as database
    original = database.SessionLocal
    database.SessionLocal = lambda: db          # the test session, not a new one
    try:
        import scripts.import_api_keys as importer
        importer.SessionLocal = lambda: db
        added = run(str(path))
    finally:
        database.SessionLocal = original
    assert added == 4

    rows = client.get("/api-keys", headers=auth("admin")).json()
    by_purpose = {r["purpose"]: r for r in rows}

    # A name that matches a tracker project becomes a real link, not loose text.
    assert by_purpose["For the imported fallback"]["project_id"] == project["id"]
    # One that does not is kept as a label - it is still work with a key.
    assert by_purpose["For the imported Scrapping"]["project_id"] is None
    assert by_purpose["For the imported Scrapping"]["project_name"] == "FD_Rate_Import"
    # "No Expiry Date" is an answer, not a missing value.
    assert by_purpose["For the imported Scrapping"]["expires_on"] is None
    assert by_purpose["For the imported Scrapping"]["expiry_state"] == "none"
    assert by_purpose["For the imported Scrapping"]["account_email"] == "fd_treasury@aikyamcap.com"
    # "Pending" in the provider column is a status, not a vendor.
    assert by_purpose["For the imported Ai agent"]["provider_id"] is None
    assert by_purpose["For the imported Ai agent"]["status"] == "pending"
    # A vendor the master list has never seen is added to it.
    assert by_purpose["For the import test mailing"]["provider_name"] == "Brand-New-Vendor"
    names = [p["name"] for p in client.get("/api-providers", headers=auth("admin")).json()]
    assert "Brand-New-Vendor" in names


def test_importing_the_same_sheet_twice_does_not_duplicate(client, seed, auth, db, tmp_path):
    """A corrected sheet gets re-imported; that must not double the register."""
    from openpyxl import Workbook
    from scripts.import_api_keys import run

    wb = Workbook()
    ws = wb.active
    ws.append(["Project", "API-Key", "Purpose", "Exp-date", "email_id"])
    ws.append(["Twice over", "Groq", "For the STT", None, None])
    path = tmp_path / "Api_Key.xlsx"
    wb.save(path)

    import scripts.import_api_keys as importer
    importer.SessionLocal = lambda: db
    assert run(str(path)) == 1
    assert run(str(path)) == 0

    rows = [r for r in client.get("/api-keys", headers=auth("admin")).json()
            if r["project_name"] == "Twice over"]
    assert len(rows) == 1
