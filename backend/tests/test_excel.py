import io
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

SAMPLE = Path(__file__).resolve().parent.parent / "data" / "Weekly_Update.xlsx"


def test_project_names_unique_case_insensitive(client, seed, auth):
    body = {"name": "Unique Name Project", "vertical_id": seed["vertical"], "status_id": seed["status_queue"]}
    assert client.post("/projects", headers=auth("admin"), json=body).status_code == 201
    r = client.post("/projects", headers=auth("admin"), json={**body, "name": "  unique name PROJECT "})
    assert r.status_code == 409
    other = client.post("/projects", headers=auth("admin"), json={**body, "name": "Rename target"}).json()
    assert client.patch(f"/projects/{other['id']}", headers=auth("admin"), json={"name": "UNIQUE NAME PROJECT"}).status_code == 409


def test_new_fields_round_trip_and_owner_ids_returned(client, seed, auth):
    body = {"name": "Fields project", "vertical_id": seed["vertical"], "status_id": seed["status_queue"],
            "assigned_by": "Vishal Trehan", "assigned_on": "2026-08-25", "remarks": "confirmation awaiting",
            "owner_ids": [seed["member"]]}
    r = client.post("/projects", headers=auth("admin"), json=body)
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["assigned_by"] == "Vishal Trehan" and p["assigned_on"] == "2026-08-25" and p["remarks"] == "confirmation awaiting"
    assert p["owner_ids"] == [seed["member"]]
    got = client.get(f"/projects/{p['id']}", headers=auth("member")).json()
    assert got["owner_ids"] == [seed["member"]]

    # A requestor sees the project - owners, dates, the lot - but never the
    # internal Remarks field.
    seen = client.get(f"/projects/{p['id']}", headers=auth("req1"))
    assert seen.status_code == 200
    assert seen.json()["owner_ids"] == [seed["member"]]
    assert seen.json()["assigned_by"] == "Vishal Trehan"
    assert seen.json()["remarks"] is None


def test_moving_to_terminal_status_stamps_completion(client, seed, auth):
    """Whether a status counts as done is a property of the status, not a
    hardcoded name - so the same move stamps or does not stamp depending on it."""
    body = {"name": "Goes live", "vertical_id": seed["vertical"], "status_id": seed["status_queue"]}
    pid = client.post("/projects", headers=auth("admin"), json=body).json()["id"]
    r = client.patch(f"/projects/{pid}", headers=auth("admin"), json={"status_id": seed["status_live"]})
    assert r.json()["actual_completion_date"] is not None

    # Turn the flag off and the same move stops meaning "finished".
    client.patch(f"/statuses/{seed['status_live']}", headers=auth("admin"), json={"is_terminal": False})
    try:
        pid2 = client.post("/projects", headers=auth("admin"), json={**body, "name": "Goes live 2"}).json()["id"]
        r = client.patch(f"/projects/{pid2}", headers=auth("admin"), json={"status_id": seed["status_live"]})
        assert r.json()["actual_completion_date"] is None
    finally:
        client.patch(f"/statuses/{seed['status_live']}", headers=auth("admin"), json={"is_terminal": True})


def test_export_has_one_sheet_per_active_status(client, seed, auth):
    assert client.get("/projects/export.xlsx", headers=auth("req1")).status_code == 403
    r = client.get("/projects/export.xlsx", headers=auth("member"))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    statuses = [s for s in client.get("/statuses", headers=auth("admin")).json() if s["is_active"] is not False]
    assert [ws.title for ws in wb.worksheets][: len(statuses)] == [s["name"] for s in sorted(statuses, key=lambda s: (s["sort_order"], s["id"]))]
    header = [c.value for c in wb.worksheets[0][1]]
    assert header[:7] == ["Project Name", "Project Description", "Assigned by", "Date of assignment",
                          "Status", "Date of completion", "Remarks"]
    names = {ws.cell(row=r, column=1).value for ws in wb.worksheets for r in range(2, ws.max_row + 1)}
    assert "Fields project" in names


def test_import_script_dry_run_and_real_import(seed, db, tmp_path):
    if not SAMPLE.exists():
        import pytest
        pytest.skip("sample workbook not present")
    import os
    env = {**os.environ}
    args = [sys.executable, "scripts/import_excel.py", str(SAMPLE), "--vertical", "Imported", "--keep-placeholders"]
    dry = subprocess.run(args + ["--dry-run"], capture_output=True, text=True, env=env, cwd=SAMPLE.parent.parent)
    assert dry.returncode == 0, dry.stdout + dry.stderr
    assert "18 project(s) would be created, 1 skipped" in dry.stdout
    from app.models.project import Project
    before = db.query(Project).count()
    real = subprocess.run(args, capture_output=True, text=True, env=env, cwd=SAMPLE.parent.parent)
    assert real.returncode == 0, real.stdout + real.stderr
    assert "18 project(s) created" in real.stdout
    db.expire_all()
    assert db.query(Project).count() == before + 18
    fdd = db.query(Project).filter(Project.name == "Financial Due Diligence").one()
    assert "confirmation awaiting from the team" in (fdd.remarks or "")
    assert fdd.assigned_by == "Kunal Rathi"
    again = subprocess.run(args, capture_output=True, text=True, env=env, cwd=SAMPLE.parent.parent)
    assert "0 project(s) created, 19 skipped" in again.stdout
