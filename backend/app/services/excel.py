"""Excel export in the team's weekly format: ONE sheet, one row per project.

    Project Name | Assigned by | Date of assignment | Status |
    <one column per calendar week> | Date of completion | Remarks

The week columns are not a stored thing. They are worked out from the calendar
at export time (see :mod:`app.services.weeks`) and filled from each project's
update history, so next week's export simply has one more column - nothing to
migrate, no job to run, nothing to keep in sync. How many weeks are shown comes
from the ``export_weeks`` setting, and where a week starts comes from the same
``digest_day_of_week`` the weekly digest uses.

The importer (``scripts/import_excel.py``) reads this exact layout back, and
shares :func:`update_to_text` / :func:`text_to_update` with it so a workbook
that goes out can come back in unchanged.
"""
import re
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app.models.project import Project
from app.models.update import Update
from app.services import daily, weeks as wk

SHEET_TITLE = "Weekly Update"

# The fixed columns, either side of however many week columns the calendar asks
# for. Widths are paired with them so the two can never fall out of step.
LEADING_COLUMNS = [("Project Name", 38), ("Assigned by", 22), ("Date of assignment", 18), ("Status", 16)]
TRAILING_COLUMNS = [("Date of completion", 18), ("Remarks", 40)]
WEEK_COLUMN_WIDTH = 34

HEADER_COLOR = "#1f2937"

# How a PPP update is written into a week cell. A plain update - progress only,
# which is nearly all of them - is written as its own text with no label, so the
# sheet reads like the one the team already keeps by hand.
FIELD_LABELS = (("Plan", "plan"), ("Progress", "progress"), ("Problem", "problem"))
_LABEL_RE = re.compile(r"^\s*(plan|progress|problem)\s*:\s*(.*)$", re.IGNORECASE)

# Two updates logged in the same week share one cell, separated by a blank line.
ENTRY_SEPARATOR = "\n\n"


def update_to_text(update: Update) -> str:
    """One update as the text that goes in its week cell."""
    filled = [(label, (getattr(update, attr) or "").strip()) for label, attr in FIELD_LABELS]
    filled = [(label, value) for label, value in filled if value]
    if not filled:
        return ""
    if len(filled) == 1 and filled[0][0] == "Progress":
        return filled[0][1]
    return "\n".join(f"{label}: {value}" for label, value in filled)


def text_to_update(text: object) -> dict[str, str | None]:
    """A week cell back into plan / progress / problem.

    Unlabelled text is progress - that is what somebody typing straight into the
    team's workbook means, and it is the exact inverse of :func:`update_to_text`.
    """
    parts: dict[str, list[str]] = {}
    current = "progress"
    for line in str(text or "").splitlines():
        match = _LABEL_RE.match(line)
        if match:
            current = match.group(1).lower()
            line = match.group(2)
        parts.setdefault(current, []).append(line)
    return {attr: ("\n".join(parts.get(attr, [])).strip() or None) for _, attr in FIELD_LABELS}


def _hex_to_fill(color: str | None) -> PatternFill:
    hex_color = (color or "#6b7280").lstrip("#")
    if len(hex_color) != 6:
        hex_color = "6b7280"
    return PatternFill("solid", fgColor=f"FF{hex_color.upper()}")


def _text_on(color: str | None) -> str:
    hex_color = (color or "#6b7280").lstrip("#")
    try:
        r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return "FFFFFFFF"
    luminance = 0.299 * r + 0.587 * g + 0.114 * b
    return "FF111827" if luminance > 160 else "FFFFFFFF"


def _local_date(moment: datetime, tz) -> date:
    """The local day an update was logged on.

    Postgres hands back an aware timestamp and SQLite a naive one; a naive value
    is UTC, which is what ``func.now()`` wrote.
    """
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(tz).date()


def _updates_by_week(db: Session, dow: int) -> dict[tuple[int, date], list[str]]:
    """Every update, filed under (project, the week it was logged in)."""
    tz = daily.tzinfo_for(db)
    grouped: dict[tuple[int, date], list[str]] = {}
    for update in db.query(Update).order_by(Update.created_at, Update.id).all():
        text = update_to_text(update)
        if not text:
            continue
        start = wk.week_start_for(_local_date(update.created_at, tz), dow)
        grouped.setdefault((update.project_id, start), []).append(text)
    return grouped


def build_projects_workbook(db: Session) -> Workbook:
    week_starts = wk.recent_weeks(db, daily.today_for(db))
    by_week = _updates_by_week(db, wk.start_dow(db))

    columns = (
        LEADING_COLUMNS
        + [(wk.header_for(start), WEEK_COLUMN_WIDTH) for start in week_starts]
        + TRAILING_COLUMNS
    )
    first_week_col = len(LEADING_COLUMNS) + 1
    remarks_col = len(columns)

    projects = (
        db.query(Project)
        .options(selectinload(Project.status))
        .order_by(Project.created_at, Project.id)
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    thin = Side(style="thin", color="FFD1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([title for title, _ in columns])
    head_fill = _hex_to_fill(HEADER_COLOR)
    head_font = Font(bold=True, color=_text_on(HEADER_COLOR))
    for col, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = ws.cell(row=2, column=first_week_col)

    for project in projects:
        ws.append([
            project.name,
            project.assigned_by,
            project.assigned_on,
            project.status.name if project.status else None,
            *[ENTRY_SEPARATOR.join(by_week.get((project.id, start), [])) or None for start in week_starts],
            project.actual_completion_date,
            project.remarks,
        ])
        row = ws.max_row
        for col in range(1, len(columns) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col >= first_week_col or col == remarks_col)
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "DD-MMM-YYYY"
        if project.status:
            status_cell = ws.cell(row=row, column=len(LEADING_COLUMNS))
            status_cell.fill = _hex_to_fill(project.status.color)
            status_cell.font = Font(color=_text_on(project.status.color))

    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 1)}"
    return wb
