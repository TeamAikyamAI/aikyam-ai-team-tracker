"""The API key register as a workbook.

Same five columns as the team's own Api_Key.xlsx, in the same order, so the
export drops straight into where that file was already used - plus the two
columns the register knows and the spreadsheet never did: the status of the key
and how many days are left on it.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# The original column names, kept exactly as the team wrote them.
COLUMNS = ["Project", "API-Key", "Purpose", "Exp-date", "email_id", "Status", "Days left"]
WIDTHS = [30, 18, 42, 16, 34, 14, 12]

_HEAD_FILL = PatternFill("solid", fgColor="FF1F2937")
_EXPIRED_FILL = PatternFill("solid", fgColor="FFFCE7E9")
_SOON_FILL = PatternFill("solid", fgColor="FFFFF4D6")
_THIN = Side(style="thin", color="FFD9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_api_key_workbook(rows, app_name: str, today: date) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "API Keys"

    ws.append(COLUMNS)
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        cell.border = _BORDER
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([
            row.project_name,
            row.provider_name,
            row.purpose or "",
            row.expires_on if row.expires_on else "No Expiry Date",
            row.account_email or "",
            row.status.capitalize(),
            "" if row.days_left is None else row.days_left,
        ])
        written = ws[ws.max_row]
        for cell in written:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row.expiry_state == "expired":
            for cell in written:
                cell.fill = _EXPIRED_FILL
        elif row.expiry_state == "soon":
            for cell in written:
                cell.fill = _SOON_FILL
        if row.expires_on:
            written[3].number_format = "dd-mmm-yyyy"

    ws.append([])
    note = ws.cell(row=ws.max_row + 1, column=1)
    note.value = (
        f"{app_name} - API key register as at {today.isoformat()}. "
        "This records which key exists and when it lapses; it never stores the key itself."
    )
    note.font = Font(italic=True, size=9, color="FF6B7280")
    return wb
